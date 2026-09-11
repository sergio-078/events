import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.utils import timezone
from datetime import datetime
import pandas as pd
from io import BytesIO

from ..models import Event


class ExcelExporter:
    """Класс для экспорта данных в Excel"""
    
    def __init__(self, queryset):
        self.queryset = queryset
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = "Заявки на работы"
        
        # Стили
        self.header_font = Font(bold=True, color="FFFFFF", size=11)
        self.header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
        self.header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.cell_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def export_to_excel(self):
        """Основной метод экспорта"""
        self._create_header()
        self._fill_data()
        self._auto_adjust_columns()
        return self._get_response()
    
    def _create_header(self):
        """Создание заголовка таблицы"""
        headers = [
            'ID', 'Статус', 'Наименование работ', 'Причина работ',
            'Начало (план)', 'Окончание (план)', 'Длительность (дн. ч. мин)',
            'Нарушение сервиса', 'Длительность перерыва',
            'Затронутые сервисы', 'Место нахождения объекта',
            'Место выполнения работ', 'Ответственный',
            'Исполнители', 'Создал', 'Дата создания'
        ]
        
        for col, header in enumerate(headers, 1):
            cell = self.ws.cell(row=1, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.border
    
    def _fill_data(self):
        """Заполнение данных"""
        status_map = dict(Event.EventStatus.choices)
        
        for row_idx, event in enumerate(self.queryset, 2):
            # ID
            self.ws.cell(row=row_idx, column=1, value=event.id)
            
            # Статус
            self.ws.cell(row=row_idx, column=2, value=status_map.get(event.status, event.status))
            
            # Наименование работ
            self.ws.cell(row=row_idx, column=3, value=event.work_name or '')
            
            # Причина работ
            self.ws.cell(row=row_idx, column=4, value=event.get_work_reason_display() or '')
            
            # Начало и окончание
            self.ws.cell(row=row_idx, column=5, value=event.planned_start.strftime('%d.%m.%Y %H:%M') if event.planned_start else '')
            self.ws.cell(row=row_idx, column=6, value=event.planned_end.strftime('%d.%m.%Y %H:%M') if event.planned_end else '')
            
            # Длительность
            duration = f"{event.planned_duration_days}дн. {event.planned_duration_hours}ч. {event.planned_duration_minutes}мин."
            self.ws.cell(row=row_idx, column=7, value=duration)
            
            # Нарушение сервиса
            self.ws.cell(row=row_idx, column=8, value='Да' if event.is_service_disruption else 'Нет')
            
            # Длительность перерыва
            if event.is_service_disruption:
                disruption = f"{event.service_disruption_duration_days}дн. {event.service_disruption_duration_hours}ч. {event.service_disruption_duration_minutes}мин."
            else:
                disruption = ''
            self.ws.cell(row=row_idx, column=9, value=disruption)
            
            # Затронутые сервисы
            services = ', '.join([s.name for s in event.affected_services.all()])
            self.ws.cell(row=row_idx, column=10, value=services)
            
            # Место нахождения объекта
            if event.is_object_outside:
                location = f"За территорией: {event.object_outside_location.name if event.object_outside_location else ''}"
            else:
                location = str(event.object_location) if event.object_location else ''
            self.ws.cell(row=row_idx, column=11, value=location)
            
            # Место выполнения работ
            if event.is_work_outside:
                work_location = f"За территорией: {event.work_outside_location.name if event.work_outside_location else ''}"
            else:
                work_location = str(event.work_location) if event.work_location else ''
            self.ws.cell(row=row_idx, column=12, value=work_location)
            
            # Ответственный
            responsible = str(event.responsible_person) if event.responsible_person else ''
            self.ws.cell(row=row_idx, column=13, value=responsible)
            
            # Исполнители
            executors = ', '.join([e.full_name for e in event.executors.all()])
            self.ws.cell(row=row_idx, column=14, value=executors)
            
            # Создал
            self.ws.cell(row=row_idx, column=15, value=event.created_by.get_full_name() if event.created_by else '')
            
            # Дата создания
            self.ws.cell(row=row_idx, column=16, value=event.created_at.strftime('%d.%m.%Y %H:%M') if event.created_at else '')
            
            # Применяем стили к ячейкам
            for col in range(1, 17):
                cell = self.ws.cell(row=row_idx, column=col)
                cell.alignment = self.cell_alignment
                cell.border = self.border
        
        # Закрепляем заголовок
        self.ws.freeze_panes = 'A2'
    
    def _auto_adjust_columns(self):
        """Автоматическая подгонка ширины колонок"""
        for col in range(1, 17):
            max_length = 0
            column_letter = get_column_letter(col)
            
            for row in range(1, self.ws.max_row + 1):
                cell_value = self.ws.cell(row=row, column=col).value
                if cell_value:
                    # Учитываем длину значения
                    length = len(str(cell_value))
                    if length > max_length:
                        max_length = length
            
            # Устанавливаем ширину с запасом
            adjusted_width = min(max_length + 2, 50)
            self.ws.column_dimensions[column_letter].width = adjusted_width
    
    def _get_response(self):
        """Получение HTTP-ответа с файлом"""
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"заявки_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        # Сохраняем в буфер
        buffer = BytesIO()
        self.wb.save(buffer)
        response.write(buffer.getvalue())
        
        return response


def export_events_to_excel(queryset):
    """Функция для экспорта заявок в Excel"""
    exporter = ExcelExporter(queryset)
    return exporter.export_to_excel()


def export_events_to_csv(queryset):
    """Экспорт в CSV для больших данных"""
    import csv
    
    response = HttpResponse(content_type='text/csv')
    filename = f"заявки_{timezone.now().strftime('%Y%m%d_%H%M%S')}.csv"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    writer = csv.writer(response, delimiter=';')
    
    # Заголовки
    headers = [
        'ID', 'Статус', 'Наименование работ', 'Причина работ',
        'Начало (план)', 'Окончание (план)', 'Длительность',
        'Нарушение сервиса', 'Длительность перерыва',
        'Затронутые сервисы', 'Место нахождения', 'Место выполнения',
        'Ответственный', 'Исполнители', 'Создал', 'Дата создания'
    ]
    writer.writerow(headers)
    
    # Данные
    status_map = dict(Event.EventStatus.choices)
    
    for event in queryset:
        row = [
            event.id,
            status_map.get(event.status, event.status),
            event.work_name or '',
            event.get_work_reason_display() or '',
            event.planned_start.strftime('%d.%m.%Y %H:%M') if event.planned_start else '',
            event.planned_end.strftime('%d.%m.%Y %H:%M') if event.planned_end else '',
            f"{event.planned_duration_days}дн. {event.planned_duration_hours}ч. {event.planned_duration_minutes}мин.",
            'Да' if event.is_service_disruption else 'Нет',
            f"{event.service_disruption_duration_days}дн. {event.service_disruption_duration_hours}ч. {event.service_disruption_duration_minutes}мин." if event.is_service_disruption else '',
            ', '.join([s.name for s in event.affected_services.all()]),
            str(event.object_location) if event.object_location else '',
            str(event.work_location) if event.work_location else '',
            str(event.responsible_person) if event.responsible_person else '',
            ', '.join([e.full_name for e in event.executors.all()]),
            event.created_by.get_full_name() if event.created_by else '',
            event.created_at.strftime('%d.%m.%Y %H:%M') if event.created_at else ''
        ]
        writer.writerow(row)
    
    return response
